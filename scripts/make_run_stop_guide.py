"""Build the operator's run/stop/collaboration guide as a .docx.

Not part of the application: a document for the people running this project about one
machine. It is generated from a script rather than hand-formatted so the content stays
reviewable in history even though the .docx does not.
"""

from __future__ import annotations

import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = pathlib.Path(__file__).resolve().parent.parent / "Run_Stop_Project.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
DANGER = RGBColor(0xA6, 0x1B, 0x1B)
CODE_BG = "F2F2F2"

# The one place these appear. Every reference in the text uses these names, so a reader who
# copies a path is copying the path that exists.
SERVER_IP = "188.245.2.148"
SSH_HOST_ALIAS = "Hetzner"
SSH_KEY = "~/.ssh/id_ed25519"
SSH_PUB = "~/.ssh/id_ed25519.pub"
SSH_CONFIG = "~/.ssh/config"
REPO_URL = "https://github.com/SergeyGer/SOWSprint.ai.git"

# The SSH client block, quoted exactly as it stands on the operator's laptops. If the real file
# changes, this string must change with it — a document that quotes a config the reader does
# not have is worse than one that quotes nothing.
SSH_CONFIG_BLOCK = f"""Host {SSH_HOST_ALIAS} {SERVER_IP}
    HostName {SERVER_IP}
    User root
    IdentityFile {SSH_KEY}
    IdentitiesOnly yes
    AddKeysToAgent yes
    ServerAliveInterval 30
    ServerAliveCountMax 4
    TCPKeepAlive yes
    PasswordAuthentication no
    KbdInteractiveAuthentication no
    PreferredAuthentications publickey"""


def heading(doc, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT


def para(doc, text: str, *, bold: bool = False, italic: bool = False, size: int = 11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p


def bullet(doc, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.add_run(text)
    return p


def numbered(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def code(doc, text: str, *, danger: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = DANGER if danger else RGBColor(0x1A, 0x1A, 0x1A)
    pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), CODE_BG)
    pr.append(shd)
    return p


def note(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED
    return p


def table(doc, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def build() -> pathlib.Path:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    title = doc.add_heading("SOWSprint.ai — запуск, остановка и совместная работа", level=0)
    for run in title.runs:
        run.font.color.rgb = ACCENT

    para(
        doc,
        "Как поднять проект, когда вы садитесь за него, как полностью остановить, когда "
        "переключаетесь на другой, и как работать над ним втроём с разных ноутбуков на "
        "одном сервере.",
        italic=True,
    )
    note(
        doc,
        "Файл лежит в корне проекта и добавлен в .gitignore — в репозиторий он не попадает "
        "и живёт только на ваших машинах.",
    )

    # ================================================================================
    heading(doc, "1. Главное, что нужно знать", 1)
    para(
        doc,
        "Контейнеры проекта настроены так, что НЕ запускаются сами при перезагрузке сервера. "
        "Это сделано намеренно: проект занимает порты, и если бы он поднимался "
        "автоматически, он мешал бы другим проектам, а вы бы искали причину в чужом коде.",
    )
    table(
        doc,
        ["Правило", "Почему"],
        [
            ("После перезагрузки сервера проект не работает", "Его нужно запустить вручную"),
            (
                "Пока проект запущен, его порты заняты",
                "Другой проект не стартует и скажет, кто держит порт",
            ),
            (
                "Dev Containers не публикуют порты на сервере",
                "Несколько разработчиков работают одновременно без конфликтов",
            ),
            ("Секреты (.env) не попадают в git", "Ключи хранятся только там, где нужны"),
        ],
    )

    heading(doc, "1.1. Что где работает", 2)
    table(
        doc,
        ["Машина", "Роль"],
        [
            (
                f"Linux-сервер (текущий реальный Linux сервер {SERVER_IP})",
                "Здесь живут Docker, контейнеры, Qdrant, DeepSeek Harness и весь код",
            ),
            (
                "Ноутбуки на Windows",
                "Только VS Code и браузер. Код и контейнеры остаются на сервере",
            ),
        ],
    )
    note(
        doc,
        "Программировать локально на ноутбуке не нужно: соберёте окружение, которое "
        "отличается от серверного, и будете искать разницу в поведении. Ноутбук — это экран "
        "и клавиатура.",
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "2. SSH-доступ с нового ноутбука", 1)
    para(
        doc,
        "Всё начинается с ключа. Пароли на сервере отключены (PasswordAuthentication no), "
        "поэтому без ключа войти не получится — и это правильно: сервер стоит в интернете, и "
        "подбор пароля по SSH идёт круглосуточно.",
    )

    heading(doc, "2.1. Создание ключа на ноутбуке", 2)
    para(doc, "Windows (PowerShell), Linux и macOS — команда одна и та же:")
    code(doc, 'ssh-keygen -t ed25519 -C "ваше-имя@ноутбук"')
    para(doc, "На вопросы отвечайте так:")
    bullet(doc, f"Путь к файлу — нажмите Enter, согласившись с {SSH_KEY}.")
    bullet(
        doc,
        "Парольная фраза (passphrase) — задайте. Ключ без пароля работает как беспарольный "
        "пропуск: украли файл — получили сервер.",
    )
    para(doc, "Проверьте, что ключ создан:")
    code(doc, f"dir {SSH_PUB}          # Windows\nls -l {SSH_PUB}       # Linux, macOS")
    note(
        doc,
        "Ключи ed25519, а не RSA: короче, быстрее и без известных слабостей. Сервер "
        "принимает именно их — в authorized_keys сейчас три ключа ed25519.",
    )

    heading(doc, "2.2. Установка открытой части на сервер", 2)
    para(
        doc,
        "На сервер отправляется только ПУБЛИЧНАЯ часть — файл с расширением .pub. Приватный "
        "ключ не покидает ноутбук никогда.",
    )
    para(doc, "Способ А — если есть доступ по паролю (только для первого ключа):")
    code(doc, f"ssh-copy-id -i {SSH_PUB} root@{SERVER_IP}")

    para(doc, "Способ Б — вручную, если ssh-copy-id недоступен (Windows без OpenSSH):")
    para(doc, "Windows:")
    code(
        doc,
        f'type {SSH_PUB} | ssh root@{SERVER_IP} "mkdir -p ~/.ssh && chmod 700 ~/.ssh && '
        f'cat >> ~/.ssh/authorized_keys && chmod 600 ~/.ssh/authorized_keys"',
    )
    para(doc, "Linux и macOS:")
    code(
        doc,
        f"cat {SSH_PUB} | ssh root@{SERVER_IP} "
        f'"mkdir -p ~/.ssh && chmod 700 ~/.ssh && cat >> ~/.ssh/authorized_keys '
        f'&& chmod 600 ~/.ssh/authorized_keys"',
    )
    para(doc, "Проверка на сервере — ключ должен быть виден:")
    code(doc, f"ssh root@{SERVER_IP} 'wc -l ~/.ssh/authorized_keys'")

    heading(doc, "2.3. Файл настроек SSH на ноутбуке", 2)
    para(
        doc,
        f"Создайте файл {SSH_CONFIG} (на Windows это C:\\Users\\ИМЯ\\.ssh\\config) и вставьте "
        f"следующее. Это ТЕКУЩЕЕ содержимое файла — именно оно работает сейчас:",
    )
    code(doc, SSH_CONFIG_BLOCK)

    para(doc, "Что здесь важно и почему:", bold=True)
    table(
        doc,
        ["Строка", "Смысл"],
        [
            (
                f"Host {SSH_HOST_ALIAS} {SERVER_IP}",
                f"Позволяет писать «ssh {SSH_HOST_ALIAS}» вместо полного адреса",
            ),
            (f"IdentityFile {SSH_KEY}", "Какой именно ключ предъявлять серверу"),
            (
                "IdentitiesOnly yes",
                "Не перебирать все ключи подряд: иначе SSH предложит серверу полдюжины ключей "
                "и получит отказ после лимита попыток",
            ),
            ("AddKeysToAgent yes", "Спросить парольную фразу один раз за сессию"),
            (
                "ServerAliveInterval 30 / CountMax 4",
                "Разрывать зависший канал через 2 минуты, а не держать вечно",
            ),
            ("PasswordAuthentication no", "Клиент даже не пытается войти паролем"),
            ("PreferredAuthentications publickey", "Только ключ, никаких других методов"),
        ],
    )

    para(doc, "Права на файл (Linux и macOS — SSH откажется работать при слишком широких):",
         bold=True)
    code(
        doc,
        f"chmod 700 ~/.ssh && chmod 600 {SSH_KEY} && chmod 644 {SSH_PUB} "
        f"&& chmod 600 {SSH_CONFIG}",
    )
    note(
        doc,
        "На Windows права настраиваются через свойства файла, вкладку «Безопасность», "
        "оставив доступ только своей учётной записи.",
    )

    heading(doc, "2.4. Проверка входа", 2)
    code(doc, f"ssh {SSH_HOST_ALIAS}")
    para(doc, "Должны попасть на сервер без вопроса о пароле — только парольная фраза ключа.")
    para(doc, "Подробный вывод, если что-то не так:", bold=True)
    code(doc, f"ssh -vvv {SSH_HOST_ALIAS} 2>&1 | tail -30")

    heading(doc, "2.5. Отдельный ключ на каждый ноутбук", 2)
    para(
        doc,
        "Не копируйте один и тот же приватный ключ на разные машины. Свой ключ на каждый "
        "ноутбук означает: потеряли ноутбук — удалили одну строку из authorized_keys, "
        "остальные продолжают работать.",
    )
    code(
        doc,
        f"ssh {SSH_HOST_ALIAS}\n"
        f"nano ~/.ssh/authorized_keys   # удалить нужную строку, сохранить\n"
        f"exit",
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "3. Запуск и остановка проекта", 1)

    heading(doc, "3.1. Запуск", 2)
    para(doc, "Windows (PowerShell), на сервере или в терминале VS Code:")
    code(doc, ".\\scripts\\start.ps1")
    para(doc, "Linux:")
    code(doc, "./scripts/start.sh")
    para(
        doc,
        "Скрипт проверит Docker, убедится, что порты свободны, поднимет сервисы, дождётся "
        "готовности приложения и напечатает адрес.",
    )
    para(doc, "Если PowerShell откажется запускать файл — один раз разрешите его:")
    code(doc, "Unblock-File .\\scripts\\start.ps1")

    heading(doc, "3.2. Остановка", 2)
    para(
        doc,
        "Останавливайте перед переходом на другой проект. Это главное действие, если на "
        "машине живёт больше одного проекта.",
    )
    code(doc, "./scripts/stop.sh                  Linux\n.\\scripts\\stop.ps1                Windows")
    para(
        doc,
        "Скрипт останавливает ТОЛЬКО контейнеры этого проекта. Чужие не трогает и в конце "
        "сообщает, какие порты остались заняты и кем.",
    )
    note(
        doc,
        "Индекс Qdrant и сгенерированные документы сохраняются. Следующий запуск занимает "
        "секунды, а не повторный разбор всей базы знаний.",
    )

    heading(doc, "3.3. Проверка состояния", 2)
    code(doc, "./scripts/status.sh               Linux\n.\\scripts\\status.ps1             Windows")
    para(
        doc,
        "Скрипт покажет контейнеры проекта, кто занимает порты и какие ещё проекты работают "
        "на машине — включая dev-контейнеры коллег.",
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "4. Dev Containers: работа с кодом на сервере", 1)
    para(
        doc,
        "Расширение Dev Containers (Microsoft) поднимает среду разработки в контейнере НА "
        "СЕРВЕРЕ, а VS Code на ноутбуке подключается к ней. Код, тесты и приложение работают "
        "рядом с серверными данными; ноутбук только рисует окно.",
    )

    heading(doc, "4.1. Как это устроено", 2)
    table(
        doc,
        ["Где", "Что происходит"],
        [
            ("Ноутбук (Windows)", "VS Code, расширение Dev Containers, SSH-ключ"),
            (
                f"Сервер (текущий реальный Linux сервер {SERVER_IP})",
                "Контейнер с Python 3.11, pytest, ruff, git, ffmpeg; рядом Qdrant",
            ),
            (
                "Браузер на ноутбуке",
                "Приложение на localhost:8000 — VS Code пробрасывает порт через SSH",
            ),
        ],
    )

    heading(doc, "4.2. Первое подключение", 2)
    numbered(doc, "Установите расширения Dev Containers и Remote - SSH (оба от Microsoft).")
    numbered(doc, f"F1 → Remote-SSH: Connect to Host → выберите {SSH_HOST_ALIAS}.")
    numbered(doc, "Откроется окно на сервере. File → Open Folder → папка проекта.")
    numbered(
        doc,
        "F1 → Dev Containers: Reopen in Container. Первый раз несколько минут — собирается "
        "образ. Дальше секунды.",
    )
    note(doc, "Папка .devcontainer уже лежит в проекте. Готовить ничего не нужно.")

    heading(doc, "4.3. Порты: почему ничего не публикуется на сервере", 2)
    para(
        doc,
        "Контейнер не занимает порты на сервере вообще. Вместо этого в devcontainer.json "
        "объявлены forwardPorts, и VS Code открывает туннель от контейнера до вашего ноутбука.",
    )
    table(
        doc,
        ["Порт", "Что это", "Где открывать"],
        [
            ("8000", "Приложение SOWSprint.ai", "http://localhost:8000 на ноутбуке"),
            ("6333", "Панель Qdrant", "http://localhost:6333/dashboard на ноутбуке"),
            ("6334", "Qdrant gRPC", "Не открывать в браузере"),
        ],
    )
    note(
        doc,
        "Именно поэтому трое разработчиков могут работать одновременно: каждый получает "
        "приложение на своём localhost, а на сервере не занято ничего.",
    )

    heading(doc, "4.4. Несколько разработчиков на одном сервере", 2)
    para(
        doc,
        "Каждый получает собственный набор: свой контейнер, свой Qdrant, свою сеть, свой том "
        "для кеша pip. Имена строятся из имени пользователя на сервере.",
    )
    code(
        doc,
        "sowsprint-dev-alice          контейнер разработчика alice\n"
        "sowsprint-dev-qdrant-alice   его база данных\n"
        "sowsprint-dev-bob            контейнер разработчика bob\n"
        "sowsprint-dev-qdrant-bob     его база данных",
    )
    para(doc, "Проверить, кто ещё работает на сервере:")
    code(doc, "./scripts/status.sh")
    code(
        doc,
        "  Other projects running on this machine\n"
        "    sowsprint-dev-alice      2 container(s)  VS Code dev container — alice\n"
        "    sowsprint-dev-bob        2 container(s)  VS Code dev container — bob",
    )
    para(
        doc,
        "Если на сервере у всех одна учётная запись (root), задайте идентификатор вручную.",
        bold=True,
    )
    para(doc, "Linux-терминал перед открытием контейнера:")
    code(doc, "export SOWSPRINT_DEV_ID=alice")
    para(doc, "Windows PowerShell:")
    code(doc, '$env:SOWSPRINT_DEV_ID = "alice"')
    note(
        doc,
        "Идентификатор должен быть одинаковым при каждом подключении, иначе VS Code поднимет "
        "новый контейнер вместо подключения к прежнему.",
    )

    heading(doc, "4.5. Что доступно внутри контейнера", 2)
    bullet(doc, "Терминал VS Code работает внутри контейнера: pytest и ruff сразу в PATH.")
    bullet(doc, "Правки в файлах применяются сразу — рабочая папка подключена, а не скопирована.")
    bullet(doc, "Доступны docker и docker compose (проброшен сокет Docker).")
    para(doc, "Запустить приложение из терминала внутри контейнера:", bold=True)
    code(doc, "python -m chainlit run app.py --host 0.0.0.0 --port 8000")
    para(doc, "Прогнать тесты и линтер:", bold=True)
    code(doc, "pytest -q\nruff check src tests app.py")

    heading(doc, "4.6. Как закрыть", 2)
    bullet(doc, "Закройте окно VS Code — контейнер остановится сам (shutdownAction: stopCompose).")
    bullet(doc, "Или явно: F1 → Dev Containers: Close Remote Connection.")
    bullet(doc, "Или из терминала на сервере:")
    code(
        doc,
        "docker compose -f .devcontainer/docker-compose.yml down          # свой\n"
        "SOWSPRINT_DEV_ID=alice docker compose -f .devcontainer/docker-compose.yml down  # чужой",
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "5. Синхронизация рабочих окружений", 1)
    para(
        doc,
        "Три вида файлов ведут себя по-разному, и путать их — источник половины проблем "
        "«у меня работает, у тебя нет».",
    )
    table(
        doc,
        ["Что", "Где хранится", "Как синхронизируется"],
        [
            ("Код, тесты, конфиги", "GitHub", "git pull / git push"),
            (".env с ключами API", "Только на сервере", "Вручную, по защищённому каналу"),
            ("Run_Stop_Project.docx", "Сервер и ноутбуки", "Вручную, scp"),
            ("Индекс Qdrant", "Том Docker на сервере", "He синхронизируется, строится заново"),
        ],
    )

    heading(doc, "5.1. Единственный источник истины — сервер", 2)
    para(
        doc,
        f"Код живёт на сервере ({SERVER_IP}) и в GitHub. Ноутбуки получают его через git. "
        f"Никогда не правьте файлы только на ноутбуке: правка не попадёт на сервер и "
        f"разойдётся с ним молча.",
    )

    heading(doc, "5.2. Как клонировать репозиторий с сервера на ноутбук", 2)
    para(doc, "Вариант А — с GitHub (обычный путь):")
    code(doc, f"git clone {REPO_URL}\ncd SOWSprint.ai")

    para(
        doc,
        "Вариант Б — прямо с сервера, если нужны и те файлы, которых нет в GitHub "
        "(.env, Run_Stop_Project.docx):",
    )
    code(
        doc,
        f"scp {SSH_HOST_ALIAS}:/root/projects/SOWSprint/.env ./SOWSprint.ai/.env\n"
        f"scp {SSH_HOST_ALIAS}:/root/projects/SOWSprint/Run_Stop_Project.docx ./SOWSprint.ai/",
    )
    note(
        doc,
        "Вариант Б нужен редко. Обычно достаточно клонировать с GitHub и скопировать только "
        "то, чего в репозитории нет.",
    )

    heading(doc, "5.3. Перенос .env: только защищённым каналом", 2)
    para(
        doc,
        ".env содержит ключи API. Он не в git и не должен пересылаться почтой, мессенджером "
        "или через облачный диск.",
    )
    para(doc, "С сервера на ноутбук:", bold=True)
    code(doc, f"scp {SSH_HOST_ALIAS}:/root/projects/SOWSprint/.env .")
    para(doc, "С ноутбука на сервер:", bold=True)
    code(
        doc,
        f"scp .env {SSH_HOST_ALIAS}:/root/projects/SOWSprint/.env\n"
        f'ssh {SSH_HOST_ALIAS} "chmod 600 /root/projects/SOWSprint/.env"',
    )
    para(doc, "Права должны быть 600 — читает только владелец:", bold=True)
    code(doc, f"ssh {SSH_HOST_ALIAS} 'stat -c \"%a %n\" /root/projects/SOWSprint/.env'")

    heading(doc, "5.4. Чем окружения могут разойтись", 2)
    table(
        doc,
        ["Причина", "Как обнаружить", "Как исправить"],
        [
            (
                "Разные версии зависимостей",
                "pip list в контейнере и на сервере",
                "Работать только в контейнере — версии заданы образом",
            ),
            ("Разный .env", "Приложение пишет, что ключ не найден", "Сравнить: diff на двух машинах"),
            (
                "Разный индекс Qdrant",
                "./scripts/status.sh",
                "Индекс строится из корпуса сам, расхождение исчезнет",
            ),
            ("Разные коммиты", "git log --oneline -1", "git pull на сервере, затем пересобрать"),
        ],
    )
    para(
        doc,
        "Правило, снимающее почти все расхождения: код запускается только внутри контейнера. "
        "Тогда версия Python, набор пакетов и системные библиотеки одинаковы у всех по "
        "определению.",
        bold=True,
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "6. DeepSeek Harness: как подключиться", 1)
    para(
        doc,
        "DeepSeek Harness (DSH) — веб-интерфейс агента. Он запущен на сервере и слушает "
        "только localhost:3080. Снаружи он недоступен, и это сделано намеренно: у него нет "
        "пароля, поэтому открыть его в интернет означало бы отдать управление сервером "
        "любому, кто угадает адрес.",
    )
    code(doc, "dsh web --host 127.0.0.1 --port 3080 --no-open")

    heading(doc, "6.1. Способ А — SSH-туннель (рекомендуется)", 2)
    para(doc, "На ноутбуке, в отдельном терминале:")
    code(doc, f"ssh -L 3080:127.0.0.1:3080 -N {SSH_HOST_ALIAS}")
    para(doc, "Затем откройте в браузере: http://127.0.0.1:3080", bold=True)
    note(
        doc,
        "-N означает «не выполнять команду, только пробросить порт». Окно терминала должно "
        "оставаться открытым. Ctrl+C закрывает туннель.",
    )

    heading(doc, "6.2. Способ Б — через VS Code, без отдельного терминала", 2)
    numbered(doc, f"F1 → Remote-SSH: Connect to Host → {SSH_HOST_ALIAS}")
    numbered(doc, "В нижней панели откройте вкладку PORTS.")
    numbered(doc, "Forward a Port → введите 3080.")
    numbered(doc, "VS Code покажет ссылку — нажмите её, откроется браузер на ноутбуке.")
    note(
        doc,
        "VS Code делает то же, что и команда ssh -L, но туннель живёт, пока открыто окно, а "
        "не отдельный терминал.",
    )

    heading(doc, "6.3. Способ В — постоянный туннель в конфиге", 2)
    para(doc, f"Если нужен всегда, добавьте в {SSH_CONFIG} отдельный блок рядом с основным:")
    code(
        doc,
        f"Host {SSH_HOST_ALIAS}-dsh\n"
        f"    HostName {SERVER_IP}\n"
        f"    User root\n"
        f"    IdentityFile {SSH_KEY}\n"
        f"    IdentitiesOnly yes\n"
        f"    LocalForward 3080 127.0.0.1:3080",
    )
    para(doc, "Затем просто:")
    code(doc, f"ssh -N {SSH_HOST_ALIAS}-dsh")

    heading(doc, "6.4. Чего делать нельзя", 2)
    bullet(doc, "Менять --host на 0.0.0.0 и открывать порт 3080 в firewall — у DSH нет аутентификации.")
    bullet(doc, "Оставлять туннель открытым на чужом или общем ноутбуке.")
    bullet(doc, "Передавать порт дальше по сети флагом -g.")

    # ================================================================================
    doc.add_page_break()
    heading(doc, "7. Один разработчик с нескольких ноутбуков", 1)
    para(
        doc,
        "Ситуация: домашний ноутбук и рабочий, оба Windows, оба заходят на один сервер. "
        "Ничего специального не нужно — но есть места, где легко потерять время.",
    )

    heading(doc, "7.1. Что делать на каждом новом ноутбуке", 2)
    numbered(doc, "Создать свой ключ SSH (раздел 2.1) и положить открытую часть на сервер.")
    numbered(doc, f"Добавить блок {SSH_HOST_ALIAS} в {SSH_CONFIG} (раздел 2.3).")
    numbered(doc, "Установить VS Code и два расширения: Remote - SSH, Dev Containers.")
    numbered(doc, "Подключиться к серверу и открыть в контейнере папку проекта.")
    note(
        doc,
        "Больше ничего копировать не нужно: код и зависимости живут на сервере. Ноутбук можно "
        "потерять без последствий, кроме своего ключа.",
    )

    heading(doc, "7.2. Что НЕ нужно копировать", 2)
    bullet(
        doc,
        ".env — он один на всех, лежит на сервере. Второй экземпляр на ноутбуке разойдётся с "
        "серверным, и вы будете искать, почему ключ «не тот».",
    )
    bullet(doc, "Код проекта — он на сервере и в GitHub.")
    bullet(doc, "Индекс Qdrant — строится на сервере.")
    bullet(doc, "Docker — на ноутбуке он не нужен вовсе.")

    heading(doc, "7.3. Идентификатор разработчика", 2)
    para(
        doc,
        "Если вы задаёте SOWSPRINT_DEV_ID вручную, используйте ОДНО и то же значение на обоих "
        "ноутбуках — тогда оба подключаются к одному контейнеру, и второй открывается за "
        "секунды.",
    )
    code(
        doc,
        'export SOWSPRINT_DEV_ID=sergey      # Linux\n$env:SOWSPRINT_DEV_ID = "sergey"    # Windows',
    )
    note(
        doc,
        "Одновременно с двух ноутбуков работать в одном контейнере можно, но правки будут "
        "пересекаться без предупреждения. Практичнее закрыть окно на одном, прежде чем "
        "открыть на другом.",
    )

    heading(doc, "7.4. Если нужно перенести незакоммиченные правки", 2)
    para(doc, "Проще всего через git, даже черновиком:")
    code(
        doc,
        "git switch -c wip/home-laptop\n"
        'git add -A && git commit -m "wip" && git push -u origin wip/home-laptop',
    )
    para(doc, "На другом ноутбуке:")
    code(doc, "git fetch && git switch wip/home-laptop")
    note(
        doc,
        "Не заканчивайте работу с незакоммиченными правками на сервере: их не видно ни на "
        "втором ноутбуке, ни в истории, и легко потерять при пересборке.",
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "8. Безопасность", 1)
    para(
        doc,
        "Сервер стоит в интернете, и попытки подбора идут постоянно. Ниже — что уже сделано "
        "и что стоит проверить.",
    )

    heading(doc, "8.1. Что уже сделано", 2)
    table(
        doc,
        ["Мера", "Состояние"],
        [
            ("Вход по паролю SSH", "Отключён (PasswordAuthentication no)"),
            ("Вход под root", "Разрешён только по ключу"),
            ("fail2ban", "Установлен и активен, тюрьма sshd"),
            ("Firewall", "Активен (ufw), политика по умолчанию — DROP"),
            ("DeepSeek Harness", "Слушает только 127.0.0.1 — снаружи недоступен"),
            ("Ключи в authorized_keys", "3 ключа ed25519"),
        ],
    )

    heading(doc, "8.2. fail2ban: проверить и настроить", 2)
    para(doc, "Состояние и статистика:")
    code(doc, "fail2ban-client status\nfail2ban-client status sshd")
    para(doc, "Вывод покажет, сколько адресов заблокировано и сколько попыток было.")
    para(doc, "Действующая конфигурация тюрьмы:", bold=True)
    code(
        doc,
        "cat /etc/fail2ban/jail.d/sshd.local 2>/dev/null "
        "|| cat /etc/fail2ban/jail.local 2>/dev/null",
    )
    para(doc, "Разумные значения для этого сервера:", bold=True)
    code(
        doc,
        "[sshd]\n"
        "enabled  = true\n"
        "port     = ssh\n"
        "maxretry = 3\n"
        "findtime = 10m\n"
        "bantime  = 1h\n"
        "bantime.increment = true\n"
        "bantime.factor    = 2\n"
        "bantime.maxtime   = 1w",
    )
    note(
        doc,
        "bantime.increment означает, что повторный нарушитель получает срок вдвое больше "
        "предыдущего — до недели. Это отсекает упорных, не требуя вашего внимания.",
    )
    para(doc, "Применить и проверить:", bold=True)
    code(doc, "systemctl restart fail2ban\nfail2ban-client status sshd")
    para(doc, "Разблокировать себя, если заблокировались (с консоли хостера):", bold=True)
    code(doc, "fail2ban-client set sshd unbanip ВАШ.IP.АДРЕС")

    heading(doc, "8.3. Firewall", 2)
    para(
        doc,
        "Открытыми должны быть только 22 (SSH), 80 и 443 (если есть веб-сервер). "
        "Остальное закрыто:",
    )
    code(
        doc,
        "ufw status numbered\nufw default deny incoming\nufw allow 22/tcp\nufw enable",
    )
    note(
        doc,
        "Порты 8000, 6333, 6334 открывать в интернет НЕ нужно: приложение доступно через "
        "SSH-туннель или Dev Containers, а Qdrant не имеет аутентификации вовсе.",
    )

    heading(doc, "8.4. Ключи и доступ", 2)
    bullet(doc, "Один ключ — один ноутбук. Потеряли ноутбук — удалили строку, остальные работают.")
    bullet(doc, "Парольная фраза на ключе обязательна: файл ключа без неё — пропуск без пароля.")
    bullet(doc, "Не пересылайте приватный ключ никуда. Он не нужен ни на сервере, ни в git.")
    bullet(doc, "Проверяйте, кто имеет доступ, раз в несколько месяцев:")
    code(doc, f"ssh {SSH_HOST_ALIAS} 'cut -c1-60 ~/.ssh/authorized_keys'")

    heading(doc, "8.5. Секреты в .env", 2)
    bullet(doc, "Проверьте, что .env не попадёт в git: git check-ignore -v .env")
    bullet(doc, "Права 600: chmod 600 .env")
    bullet(doc, "Не копируйте .env в облачные диски и мессенджеры.")
    bullet(
        doc,
        "Ключи API отзываются и перевыпускаются без переустановки: это одна строка в .env и "
        "перезапуск контейнера.",
    )

    heading(doc, "8.6. Docker и сокет Docker", 2)
    para(
        doc,
        "Dev-контейнеру проброшен /var/run/docker.sock. Это даёт ему права root на сервере — "
        "так и задумано, чтобы из терминала работали docker и docker compose.",
        bold=True,
    )
    bullet(doc, "Не подключайте к этому сокету чужие контейнеры и образы.")
    bullet(doc, "Не запускайте в контейнере код, которому не доверяете.")
    bullet(doc, "Если проброс сокета не нужен — удалите строку из .devcontainer/docker-compose.yml.")

    heading(doc, "8.7. Регулярно", 2)
    code(
        doc,
        "apt update && apt upgrade -y          # обновления системы\n"
        "docker system df                     # сколько занимают образы\n"
        "docker image prune -f                # убрать неиспользуемые образы\n"
        "du -sh /var/lib/docker               # сколько занимает Docker целиком",
    )

    # ================================================================================
    doc.add_page_break()
    heading(doc, "9. Если что-то не работает", 1)
    table(
        doc,
        ["Сообщение", "Что делать"],
        [
            (
                "port 8000 is held by container 'X' of project 'Y'",
                "Порт занят другим проектом. Остановите его или выполните ./scripts/status.sh.",
            ),
            (
                "the dev container is running — that is the VS Code window",
                "Запущен dev-контейнер. Закройте окно VS Code или выполните "
                "docker compose -f .devcontainer/docker-compose.yml down.",
            ),
            (
                "Conflict. The container name ... is already in use",
                "Кто-то работает под тем же SOWSPRINT_DEV_ID. Задайте свой: "
                "export SOWSPRINT_DEV_ID=ваше-имя",
            ),
            (
                "Docker is installed but not responding",
                "Linux: sudo systemctl start docker. Windows: запустите Docker Desktop.",
            ),
            (
                "Permission denied (publickey)",
                "Ключ не принят. Проверьте: ssh -vvv Hetzner. Убедитесь, что открытая часть "
                "добавлена в ~/.ssh/authorized_keys на сервере.",
            ),
            ("Скрипт не запускается в PowerShell", "Unblock-File .\\scripts\\start.ps1"),
            (
                "Файлы не сохраняются в VS Code",
                "Проблема с владельцем файлов. Dev Containers: Rebuild Container.",
            ),
            (
                "Приложение отвечает, но без ключей API",
                "Нет .env или он пуст. Проверьте: ls -l .env на сервере.",
            ),
        ],
    )

    heading(doc, "Где что лежит", 1)
    for path, purpose in (
        ("docker-compose.yml", 'Рабочий стек. restart: "no" — не стартует сам.'),
        ("scripts/start.sh, start.ps1", "Запуск с проверкой портов."),
        ("scripts/stop.sh, stop.ps1", "Остановка только этого проекта."),
        ("scripts/status.sh, status.ps1", "Состояние контейнеров и портов."),
        ("scripts/_common.sh, _common.ps1", "Общая логика скриптов."),
        (".devcontainer/", "Настройка VS Code Dev Containers."),
        (".env", "Ключи API. В репозиторий не попадает."),
        ("Run_Stop_Project.docx", "Этот документ. В репозиторий не попадает."),
    ):
        p = doc.add_paragraph()
        run = p.add_run(path)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        p.add_run(" — " + purpose)

    doc.add_paragraph()
    note(
        doc,
        "Остановка не удаляет данные. Полное удаление выполняется отдельно и осознанно: "
        "docker compose down -v",
    )
    note(doc, f"Сервер: {SERVER_IP} · хост в SSH-конфиге: {SSH_HOST_ALIAS} · ключ: {SSH_KEY}")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"  ✅ {path.name} — {path.stat().st_size / 1024:.0f} KB")
